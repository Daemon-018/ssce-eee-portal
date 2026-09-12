"""marks_parser.py — Extract marks from Excel/PDF files + AI structuring.

Flow:
  1. Extract raw table/text from file (openpyxl for Excel, PyPDF2 for PDF)
  2. Send extracted text to OpenRouter AI for structuring
  3. Return list of {username, subject, exam, exam_marks, assign_marks}
"""
import json
import os
import re
from pathlib import Path

# Load .env if present (never crash when missing — e.g. Render uses env vars)
_env = {}
_env_file = Path(__file__).parent.joinpath(".env")
if _env_file.exists():
    for line in _env_file.read_text().splitlines():
        if "=" in line and not line.startswith("#"):
            k, v = line.split("=", 1)
            _env[k.strip()] = v.strip()

OPENROUTER_KEY = os.environ.get("OPENROUTER_API_KEY", "") or _env.get("OPENROUTER_API_KEY", "")
OPENROUTER_URL = "https://openrouter.ai/api/v1/chat/completions"
# Use a reliable model for file parsing — llama-3.3-70b handles structured extraction well
AI_MODEL = "meta-llama/llama-3.3-70b-instruct"


def extract_excel(filepath: str) -> str:
    """Extract all data from Excel as CSV-like text."""
    import openpyxl
    wb = openpyxl.load_workbook(filepath, data_only=True)
    lines = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"=== Sheet: {sheet_name} ===")
        for row in ws.iter_rows(values_only=True):
            # Skip fully empty rows
            vals = [str(c) if c is not None else "" for c in row]
            if any(v.strip() for v in vals):
                lines.append(" | ".join(vals))
    return "\n".join(lines)


def extract_pdf(filepath: str) -> str:
    """Extract text from PDF."""
    from PyPDF2 import PdfReader
    reader = PdfReader(filepath)
    lines = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            lines.append(text)
    return "\n".join(lines)


EXTRACTORS = {
    ".xlsx": extract_excel,
    ".xls": extract_excel,
    ".pdf": extract_pdf,
}


def parse_marks_file(filepath: str, available_subjects: list[str] = None) -> list[dict]:
    """Parse a marks file and return structured records via AI.

    Returns list of:
      {username, subject, exam, exam_marks, assign_marks}

    Raises ValueError on parse failure.
    """
    ext = Path(filepath).suffix.lower()
    if ext not in EXTRACTORS:
        raise ValueError(f"Unsupported file type: {ext}. Use .xlsx or .pdf")

    raw_text = EXTRACTORS[ext](filepath)
    if not raw_text.strip():
        raise ValueError("File appears empty or unreadable.")

    if not OPENROUTER_KEY:
        raise ValueError("OpenRouter API key not configured.")

    # Build subject list for the AI prompt
    subj_hint = ""
    if available_subjects:
        subj_hint = (
            "\n\nKnown subjects in the system (match these exactly):\n"
            + "\n".join(f"  - {s}" for s in available_subjects)
        )

    prompt = f"""You are a marks data extraction assistant for an EEE college department.

A faculty member uploaded a marks sheet. Extract ALL marks records from it.

The extracted file content is below between ---FILE START--- and ---FILE END---.

---FILE START---
{raw_text}
---FILE END---
{subj_hint}

RULES:
1. Each student row becomes ONE record per subject.
2. "exam" field is ALWAYS "MID1" or "MID2" (never "MID 1" with space, never "RESULT").
3. "exam_marks" is the exam/test score (max 25). If only one mark column, assume exam_marks.
4. "assign_marks" is the assignment score (max 5). If no assignment column, default to 0.
5. "username" is the student roll number EXACTLY as in the file (e.g. "24W61A0201").
6. Match subject names to the known subjects list as closely as possible.
7. Skip header rows, summary rows, total rows, pass/fail rows.
8. If marks column says a number out of range (>25 for exam, >5 for assignment), cap it.
9. If you see "A" or "absent" for a student's mark, use 0.

Return ONLY a JSON array (no markdown, no explanation):
[
  {{"username": "24W61A0201", "subject": "Power Electronics", "exam": "MID1", "exam_marks": 18, "assign_marks": 4}},
  ...
]

Return empty array [] if the file has no parseable marks data."""

    import urllib.request

    payload = json.dumps({
        "model": AI_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0.0,
        "max_tokens": 8000,
    }).encode()

    req = urllib.request.Request(
        OPENROUTER_URL,
        data=payload,
        headers={
            "Authorization": f"Bearer {OPENROUTER_KEY}",
            "Content-Type": "application/json",
            "HTTP-Referer": "https://ssce-eee-department.onrender.com",
        },
    )

    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            data = json.loads(resp.read())
    except Exception as e:
        raise ValueError(f"AI API call failed: {e}")

    content = data.get("choices", [{}])[0].get("message", {}).get("content", "")

    # Extract JSON array from response (handle markdown code blocks)
    content = content.strip()
    if content.startswith("```"):
        # Strip markdown code fence
        content = re.sub(r"^```(?:json)?\s*\n?", "", content)
        content = re.sub(r"\n?```\s*$", "", content)

    try:
        records = json.loads(content)
    except json.JSONDecodeError:
        # Try to find JSON array in the text
        match = re.search(r"\[.*\]", content, re.DOTALL)
        if match:
            records = json.loads(match.group())
        else:
            raise ValueError(f"AI returned invalid JSON: {content[:200]}")

    if not isinstance(records, list):
        raise ValueError("AI returned non-array result.")

    # Validate each record
    def _to_int(v, lo, hi):
        try:
            n = int(float(str(v).strip() or 0))
        except (TypeError, ValueError):
            return lo
        return max(lo, min(hi, n))

    valid = []
    for r in records:
        if not all(k in r for k in ("username", "subject", "exam")):
            continue
        r["username"] = str(r["username"]).strip()
        r["subject"] = str(r["subject"]).strip()
        r["exam"] = str(r["exam"]).upper().replace(" ", "")
        if r["exam"] not in ("MID1", "MID2"):
            r["exam"] = "MID1"
        r["exam_marks"] = _to_int(r.get("exam_marks", 0), 0, 25)
        r["assign_marks"] = _to_int(r.get("assign_marks", 0), 0, 5)
        if r["username"] and r["subject"]:
            valid.append(r)

    return valid
