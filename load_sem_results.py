#!/usr/bin/env python3
"""Parse 1-1 semester result Excel files and load into sem_results table."""
import openpyxl
import sqlite3
import os
from pathlib import Path

DB_PATH = Path(__file__).parent / "eee.db"
EXCEL_DIR = Path("/sdcard/Download/college website")

# Map filenames to exam cycle and label
EXAM_MAP = {
    "1-1 DEC - 2024 (1) .xlsx": (1, "DEC 2024"),
    "1-1 JUN 2025 (2 ) .xlsx": (2, "JUN 2025"),
    "1-1 DEC - 2025 (3) .xlsx": (3, "DEC 2025"),
    "1-1 JUN 2026 (4) .xlsx": (4, "JUN 2026"),
}

YEAR_SEM = "1-1"

def parse_excel(filepath, exam_cycle, exam_label):
    """Parse one Excel file → list of dicts."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    records = []
    
    # Parse Main sheet for SGPA summary
    main = wb["Main"]
    main_rows = list(main.iter_rows(min_row=2, values_only=True))
    sgpa_map = {}
    for row in main_rows:
        if not row[2]:  # skip empty
            continue
        ht = str(row[2]).strip()
        sgpa = row[5] if row[5] else 0
        total_credit = int(row[6]) if row[6] else 0
        total_appeared = int(row[7]) if row[7] else 0
        total_passed = int(row[8]) if row[8] else 0
        sgpa_map[ht] = {
            "sgpa": float(sgpa) if sgpa else 0,
            "total_credit": total_credit,
            "total_appeared": total_appeared,
            "total_passed": total_passed,
        }
    
    # Parse subject sheets (S2-S10)
    for sheet_name in wb.sheetnames:
        if not sheet_name.startswith("S"):
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        for row in rows:
            if not row[2]:  # skip empty rows
                continue
            ht = str(row[2]).strip()
            subject_name = str(row[3]).strip() if row[3] else ""
            code = str(row[4]).strip() if row[4] else ""
            grade = str(row[5]).strip() if row[5] else ""
            grade_point = int(row[6]) if row[6] else 0
            credit = int(row[7]) if row[7] else 0
            credit_status = str(row[8]).strip() if row[8] else ""
            
            # Get SGPA from main sheet
            sgpa_info = sgpa_map.get(ht, {})
            
            records.append({
                "username": ht,
                "year_sem": YEAR_SEM,
                "subject": subject_name,
                "code": code,
                "grade": grade,
                "grade_point": grade_point,
                "credit": credit,
                "credit_status": credit_status,
                "exam_cycle": exam_cycle,
                "exam_label": exam_label,
                "sgpa": sgpa_info.get("sgpa", 0),
                "total_credits": sgpa_info.get("total_credit", 0),
                "total_appeared": sgpa_info.get("total_appeared", 0),
                "total_passed": sgpa_info.get("total_passed", 0),
            })
    
    wb.close()
    return records

def load_to_db(records):
    """Insert records into sem_results table."""
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    
    inserted = 0
    updated = 0
    for r in records:
        try:
            cur.execute("""
                INSERT INTO sem_results 
                (username, year_sem, subject, code, grade, grade_point, credit, 
                 credit_status, exam_cycle, exam_label, sgpa, total_credits, 
                 total_appeared, total_passed)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                r["username"], r["year_sem"], r["subject"], r["code"],
                r["grade"], r["grade_point"], r["credit"], r["credit_status"],
                r["exam_cycle"], r["exam_label"], r["sgpa"], r["total_credits"],
                r["total_appeared"], r["total_passed"]
            ))
            inserted += 1
        except sqlite3.IntegrityError:
            # Update existing record
            cur.execute("""
                UPDATE sem_results SET 
                grade=?, grade_point=?, credit=?, credit_status=?, 
                sgpa=?, total_credits=?, total_appeared=?, total_passed=?,
                exam_label=?, updated_at=datetime('now','localtime')
                WHERE username=? AND year_sem=? AND subject=? AND exam_cycle=?
            """, (
                r["grade"], r["grade_point"], r["credit"], r["credit_status"],
                r["sgpa"], r["total_credits"], r["total_appeared"], r["total_passed"],
                r["exam_label"], r["username"], r["year_sem"], r["subject"], r["exam_cycle"]
            ))
            updated += 1
    
    conn.commit()
    conn.close()
    return inserted, updated

def main():
    print("=== Loading 1-1 semester results ===")
    
    # Ensure table exists
    conn = sqlite3.connect(DB_PATH)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS sem_results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL,
            year_sem TEXT NOT NULL,
            subject TEXT NOT NULL,
            code TEXT DEFAULT '',
            grade TEXT DEFAULT '',
            grade_point INTEGER DEFAULT 0,
            credit INTEGER DEFAULT 0,
            credit_status TEXT DEFAULT '',
            exam_cycle INTEGER DEFAULT 1,
            exam_label TEXT DEFAULT '',
            sgpa REAL DEFAULT 0,
            total_credits INTEGER DEFAULT 0,
            total_appeared INTEGER DEFAULT 0,
            total_passed INTEGER DEFAULT 0,
            updated_at TEXT DEFAULT (datetime('now','localtime')),
            UNIQUE(username, year_sem, subject, exam_cycle)
        )
    """)
    conn.commit()
    conn.close()
    
    total_inserted = 0
    total_updated = 0
    
    for filename, (cycle, label) in EXAM_MAP.items():
        filepath = EXCEL_DIR / filename
        if not filepath.exists():
            print(f"SKIP: {filename} not found")
            continue
        
        print(f"\nProcessing: {filename}")
        records = parse_excel(filepath, cycle, label)
        print(f"  Found {len(records)} subject records")
        
        inserted, updated = load_to_db(records)
        print(f"  Inserted: {inserted}, Updated: {updated}")
        total_inserted += inserted
        total_updated += updated
    
    print(f"\n=== Done ===")
    print(f"Total inserted: {total_inserted}")
    print(f"Total updated: {total_updated}")
    
    # Verify
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    count = conn.execute("SELECT COUNT(*) FROM sem_results").fetchone()[0]
    students = conn.execute("SELECT COUNT(DISTINCT username) FROM sem_results").fetchone()[0]
    print(f"DB now has {count} records for {students} students")
    conn.close()

if __name__ == "__main__":
    main()
